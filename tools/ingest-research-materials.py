#!/usr/bin/env python3
"""Read-only local research corpus indexing and referenced text extraction.

No Office automation, macros, remote fetches or source-file rewrites. Extracted
text is source evidence, never instructions or a claim of official applicability.
"""
from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
import os
import posixpath
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import sys
import time
import zipfile
from xml.etree import ElementTree as ET

FOLDERS = (
    '05、15章市场调查理论与方法电子书籍',
    '09、17种类别调查问卷案例及参考资料',
    '08、【报告撰写】撰写报告避坑意见资料及评分标准',
    '03、10章决赛市场调查理论与方法ppt',
    '01、50套历年正大杯市场调研大赛网部分真题和模拟题考题库',
    '10、2024年获奖作品6篇+8篇PPT',
)
MAX_SOURCE_BYTES = 300_000_000
MAX_MEMBER_BYTES = 100_000_000
MAX_PACKAGE_BYTES = 500_000_000
MAX_MEMBERS = 20_000
MAX_CELLS = 2_000_000
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
OLE = bytes.fromhex('d0cf11e0a1b11ae1')


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def atomic_json(path, value):
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def safe_name(name):
    normalized = name.replace('\\', '/')
    pure = PurePosixPath(normalized)
    return bool(normalized and not normalized.startswith('/') and not re.match(r'^[A-Za-z]:', normalized)
                and '..' not in pure.parts and not any(ord(c) < 32 for c in normalized))


def zip_inventory(path):
    entries, warnings = [], []
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_MEMBERS:
            raise ValueError('package_member_count_limit')
        for item in infos:
            mode = item.external_attr >> 16
            unsafe = not safe_name(item.filename) or (mode & 0o170000) == 0o120000
            risk = []
            if unsafe:
                risk.append('unsafe_path_or_link')
            if item.flag_bits & 1:
                risk.append('encrypted')
            if item.file_size > MAX_MEMBER_BYTES or (item.compress_size and item.file_size / item.compress_size > 2000):
                risk.append('expansion_limit')
            entries.append({'member': item.filename, 'bytes': item.file_size, 'compressed_bytes': item.compress_size,
                            'directory': item.is_dir(), 'risk': risk})
        if sum(i.file_size for i in infos) > MAX_PACKAGE_BYTES:
            warnings.append('package_expansion_limit')
    return entries, warnings


def safe_package(path):
    entries, warnings = zip_inventory(path)
    # Large videos/images are inventoried but never decompressed here. Bounds
    # apply to parts actually parsed; unsafe paths/encryption reject all parts.
    if any('unsafe_path_or_link' in item['risk'] or 'encrypted' in item['risk']
           or ('expansion_limit' in item['risk'] and item['member'].lower().endswith('.xml'))
           for item in entries):
        raise ValueError('unsafe_or_excessive_package')
    return entries


def xml_root(archive, name):
    blob = archive.read(name)
    if len(blob) > MAX_MEMBER_BYTES or b'<!DOCTYPE' in blob.upper() or b'<!ENTITY' in blob.upper():
        raise ValueError('xml_entities_or_size_rejected')
    return ET.fromstring(blob)


def paragraph_text(paragraph, word=False):
    prefix = 'w' if word else 'a'
    text = []
    for item in paragraph.iter():
        if item.tag == '{%s}t' % NS[prefix]:
            text.append(item.text or '')
        elif item.tag.endswith('}tab'):
            text.append('\t')
        elif item.tag.endswith('}br') or item.tag.endswith('}cr'):
            text.append('\n')
    return ''.join(text).strip()


def extract_docx(path):
    safe_package(path)
    records = []
    with zipfile.ZipFile(path) as archive:
        parts = [name for name in archive.namelist() if re.match(r'^word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$', name)]
        for part in sorted(parts, key=lambda name: (name != 'word/document.xml', name)):
            root = xml_root(archive, part)
            for i, paragraph in enumerate(root.iter('{%s}p' % NS['w']), 1):
                text = paragraph_text(paragraph, word=True)
                if text:
                    records.append({'ref': f'{part}#paragraph-{i}', 'text': text})
        if 'word/document.xml' not in parts:
            raise ValueError('missing_word_document')
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
    return records, {'parts': parts, 'media_files': len(media), 'unread_media': len(media)}, ['Embedded images are indexed, not OCR-read.'] if media else []


def extract_pptx(path):
    safe_package(path)
    records = []
    relns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    idns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(path) as archive:
        presentation = xml_root(archive, 'ppt/presentation.xml')
        relations = xml_root(archive, 'ppt/_rels/presentation.xml.rels')
        targets = {item.attrib['Id']: posixpath.normpath('ppt/' + item.attrib['Target'])
                   for item in relations if item.attrib.get('TargetMode') != 'External'}
        slides = [targets[item.attrib['{' + idns + '}id']]
                  for item in presentation.findall('p:sldIdLst/p:sldId', NS)]
        if not slides or any(not safe_name(name) or not name.startswith('ppt/slides/') for name in slides):
            raise ValueError('missing_or_invalid_presentation_slides')
        ordered_parts = []
        for sequence, part in enumerate(slides, 1):
            ordered_parts.append((sequence, part, 'slide'))
            relpath = posixpath.dirname(part) + '/_rels/' + posixpath.basename(part) + '.rels'
            if relpath in archive.namelist():
                for rel in xml_root(archive, relpath):
                    if rel.attrib.get('Type', '').endswith('/notesSlide') and rel.attrib.get('TargetMode') != 'External':
                        target = posixpath.normpath(posixpath.dirname(part) + '/' + rel.attrib['Target'])
                        if not safe_name(target) or not target.startswith('ppt/notesSlides/'):
                            raise ValueError('invalid_note_part')
                        ordered_parts.append((sequence, target, 'notes'))
        for sequence, part, kind in ordered_parts:
            root = xml_root(archive, part)
            for i, paragraph in enumerate(root.iter('{%s}p' % NS['a']), 1):
                text = paragraph_text(paragraph)
                if text:
                    records.append({'ref': f'slide-{sequence}/{kind}/{part}#paragraph-{i}',
                                    'slide': sequence, 'part': part, 'kind': kind, 'text': text})
        media = [name for name in archive.namelist() if name.startswith('ppt/media/')]
    return records, {'slides': len(slides), 'slide_order': slides, 'parts': [part for _, part, _ in ordered_parts], 'media_files': len(media), 'unread_media': len(media)}, ['Embedded images/video and chart/SmartArt objects are not OCR-read, played or semantically reconstructed; slide and note text only.'] if media else []


def extract_xlsx(path):
    safe_package(path)
    if not importlib.util.find_spec('openpyxl'):
        raise RuntimeError('openpyxl_unavailable')
    import openpyxl
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    records, sheets, count = [], [], 0
    try:
        for sheet in formulas:
            sheets.append(sheet.title)
            for row in sheet.iter_rows():
                for cell in row:
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError('spreadsheet_cell_limit')
                    if cell.value is None:
                        continue
                    record = {'ref': f'{sheet.title}!{cell.coordinate}', 'text': str(cell.value)}
                    if cell.data_type == 'f':
                        record['formula'] = str(cell.value)
                        record['cached_value'] = cached[sheet.title][cell.coordinate].value
                    records.append(record)
    finally:
        formulas.close()
        cached.close()
    return records, {'sheets': sheets, 'visited_cells': count}, ['Formula text and stored caches are recorded; no spreadsheet engine recalculation.']


def extract_pdf(path):
    converter = shutil.which('pdftotext')
    if converter:
        response = subprocess.run([converter, '-layout', '-enc', 'UTF-8', str(path), '-'], capture_output=True, timeout=120, check=False)
        if response.returncode == 0:
            content = response.stdout.decode('utf-8', errors='strict')
            pages = content.split('\f')
            if pages and not pages[-1].strip():
                pages.pop()
            records = [{'ref': f'page-{i}', 'text': text.strip()} for i, text in enumerate(pages, 1) if text.strip()]
            empty = [i for i, text in enumerate(pages, 1) if not text.strip()]
            return records, {'pages': len(pages), 'empty_text_pages': empty, 'engine': 'pdftotext-layout'}, ['Some pages have no extracted text; their visual content has not been read.'] if empty else []
    if not importlib.util.find_spec('pypdf'):
        raise RuntimeError('pdf_text_engines_unavailable')
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise RuntimeError('encrypted_pdf')
    records, empty = [], []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ''
        if text.strip():
            records.append({'ref': f'page-{i}', 'text': text.strip()})
        else:
            empty.append(i)
    return records, {'pages': len(reader.pages), 'empty_text_pages': empty, 'engine': 'pypdf'}, ['Some pages have no extracted text; their visual content has not been read.'] if empty else []


def extract_text(path):
    raw = path.read_bytes()
    for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
        try:
            content = raw.decode(encoding)
        except UnicodeError:
            continue
        if '\x00' in content or sum(ord(c) < 9 for c in content) > len(content) / 100:
            continue
        return [{'ref': f'line-{i}', 'text': line} for i, line in enumerate(content.splitlines(), 1) if line.strip()], {'encoding': encoding}, []
    raise RuntimeError('unknown_binary_or_text_encoding')


def read_one(path):
    with path.open('rb') as stream:
        head = stream.read(16)
    ext = path.suffix.lower()
    if head.startswith(b'PK'):
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        if 'word/document.xml' in names:
            return 'docx', extract_docx(path)
        if 'ppt/presentation.xml' in names:
            return 'pptx', extract_pptx(path)
        if 'xl/workbook.xml' in names:
            return 'xlsx', extract_xlsx(path)
        members, warnings = zip_inventory(path)
        return 'zip', ([], {'members': members, 'member_count': len(members), 'inventory_only': True}, warnings)
    if head.startswith(b'%PDF'):
        return 'pdf', extract_pdf(path)
    if head.startswith(OLE):
        raise RuntimeError('legacy_ole_converter_unavailable: no antiword/catdoc/LibreOffice/xlrd; Office automation and macros disabled')
    if ext in ('.txt', '.csv', '.md', '.html', '.htm', '.xml', ''):
        return 'text', extract_text(path)
    raise RuntimeError(f'unsupported_or_mislabeled_format:{ext}')


def merge_legacy(out):
    manifest_path = out / 'manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
    seen = {entry['id']: entry for entry in manifest['entries']}
    for receipt_path in (out / 'legacy').glob('*.receipt.json'):
        if not re.fullmatch(r'material-[0-9]{3}\.receipt\.json', receipt_path.name):
            continue  # Retained rejection/fence evidence is not the authoritative receipt.
        receipt = json.loads(receipt_path.read_text(encoding='utf-8-sig'))
        entry = seen.get(receipt.get('id'))
        if not entry or entry['sha256'] != receipt.get('sha256'):
            raise ValueError('legacy_receipt_identity_mismatch')
        status = receipt['status']
        if status == 'running':
            continue
        if status not in ('extracted', 'no_text', 'timeout', 'error'):
            raise ValueError('legacy_receipt_unknown_status')
        if receipt.get('source_unchanged') is False:
            raise ValueError('source_changed_in_legacy_receipt')
        entry['initial_status'] = entry.get('initial_status', entry['status'])
        entry['status'] = status
        entry['format'] = entry['extension'].lstrip('.')
        entry['legacy_receipt'] = receipt_path.relative_to(out).as_posix()
        entry['legacy_engine'] = receipt.get('engine')
        entry['detail'] = receipt
        if status in ('extracted', 'no_text'):
            text_path = out / 'legacy' / receipt['text_file']
            if text_path.parent != out / 'legacy' or digest(text_path) != receipt['text_sha256']:
                raise ValueError('legacy_text_hash_mismatch')
            entry['text_path'] = text_path.relative_to(out).as_posix()
            entry['text_sha256'] = receipt['text_sha256']
            entry['records'] = receipt['records']
            entry['text_chars'] = receipt['text_chars']
            entry['warnings'] = [receipt['unread_embedded_objects']] if receipt.get('unread_embedded_objects') else []
            entry.pop('error', None)
            entry.pop('error_type', None)
        elif status == 'timeout':
            entry['error'] = 'Office read-only helper timed out; text is not accepted as extracted'
        else:
            entry['error'] = receipt.get('error', 'legacy_helper_error')
    for entry in manifest['entries']:
        if entry['status'] == 'duplicate':
            original = seen[entry['duplicate_of']]
            entry.update({'content_status': original['status'], 'text_path': original.get('text_path'), 'format': original.get('format')})
    manifest['counts'] = {status: sum(entry['status'] == status for entry in manifest['entries']) for status in sorted({entry['status'] for entry in manifest['entries']})}
    manifest['legacy_updated_at_utc'] = time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
    atomic_json(manifest_path, manifest)
    (out / 'manifest.sha256').write_text(digest(manifest_path) + '  manifest.json\n', encoding='ascii')
    return manifest


def ingest_archive_members(source, out):
    manifest_path = out / 'manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
    nested, container_receipts = [], []
    allowed = {'.docx', '.pptx', '.xlsx', '.pdf', '.txt', '.csv', '.md'}
    archive_out = out / 'archive-members'
    archive_out.mkdir(exist_ok=True)
    for entry in manifest['entries']:
        if entry['status'] != 'inventory_only':
            continue
        path = source / entry['source_relative']
        hash_before = digest(path)
        if hash_before != entry['sha256']:
            raise ValueError('archive_source_hash_mismatch')
        members, warnings = zip_inventory(path)
        with zipfile.ZipFile(path) as archive:
            for item in members:
                if item['directory']:
                    continue
                name = item['member']
                record = {'archive_id': entry['id'], 'archive_sha256': entry['sha256'],
                          'member': name, 'bytes': item['bytes'], 'risk': item['risk']}
                nested.append(record)
                if item['risk'] or warnings:
                    record.update({'status': 'blocked', 'error': 'unsafe_or_oversized_archive_member'})
                    continue
                ext = PurePosixPath(name).suffix.lower()
                if ext not in allowed:
                    record.update({'status': 'blocked', 'error': 'no_safe_available_reader_for_member_type:' + ext})
                    continue
                try:
                    blob = archive.read(name)
                    member_hash = hashlib.sha256(blob).hexdigest()
                    record['sha256'] = member_hash
                    raw_path = archive_out / (member_hash + ext)
                    raw_path.write_bytes(blob)
                    detected, (records, detail, member_warnings) = read_one(raw_path)
                    text_path = archive_out / (member_hash + '.jsonl')
                    with text_path.open('w', encoding='utf-8', newline='\n') as stream:
                        for text_record in records:
                            text_record['ref'] = 'zip-member:' + name + '/' + text_record['ref']
                            stream.write(json.dumps(text_record, ensure_ascii=False, default=str) + '\n')
                    record.update({'status': 'extracted' if records else 'no_text',
                                   'format': detected, 'records': len(records),
                                   'text_path': text_path.relative_to(out).as_posix(),
                                   'text_sha256': digest(text_path), 'detail': detail, 'warnings': member_warnings})
                except Exception as error:
                    record.update({'status': 'error', 'error_type': type(error).__name__, 'error': str(error)})
        hash_after = digest(path)
        container_receipts.append({'archive_id': entry['id'], 'sha256_before': hash_before,
                                   'sha256_after': hash_after, 'source_unchanged': hash_before == hash_after})
        if hash_after != hash_before:
            raise ValueError('archive_source_changed_during_read')
        entry['archive_member_inventory'] = 'archive-members.json'
    atomic_json(out / 'archive-members.json', {'schema': 'local-archive-members/v1',
                                             'top_level_file_count_unchanged': True,
                                             'entries': nested, 'containers': container_receipts,
                                             'counts': {status: sum(item['status'] == status for item in nested) for status in sorted({item['status'] for item in nested})}})
    manifest['archive_readable_member_count'] = sum(item['status'] == 'extracted' for item in nested)
    manifest['archive_blocked_member_count'] = sum(item['status'] == 'blocked' for item in nested)
    manifest['archive_error_member_count'] = sum(item['status'] == 'error' for item in nested)
    atomic_json(manifest_path, manifest)
    (out / 'manifest.sha256').write_text(digest(manifest_path) + '  manifest.json\n', encoding='ascii')
    return nested


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--expected-files', type=int, default=283)
    parser.add_argument('--archive-members', action='store_true', help='Read safe supported members to excluded local files; never load executables.')
    parser.add_argument('--merge-legacy', action='store_true', help='Merge already written local legacy receipts; do not re-extract modern files.')
    args = parser.parse_args()
    source = args.source_root.resolve(strict=True)
    out = args.output.resolve()
    if out == source or source in out.parents:
        parser.error('Output must be outside the source tree.')
    out.mkdir(parents=True, exist_ok=True)
    (out / 'text').mkdir(exist_ok=True)
    if args.archive_members:
        nested = ingest_archive_members(source, out)
        print(json.dumps({'archive_members': len(nested), 'extracted': sum(item['status'] == 'extracted' for item in nested), 'blocked': sum(item['status'] == 'blocked' for item in nested), 'errors': sum(item['status'] == 'error' for item in nested)}, ensure_ascii=False))
        return 0
    if args.merge_legacy:
        manifest = merge_legacy(out)
        print(json.dumps({'legacy_merged': True, 'files': manifest['files'], 'unique_hashes': manifest['unique_hashes'], 'counts': manifest['counts']}, ensure_ascii=False))
        return 0
    entries, by_hash = [], {}
    paths = []
    for folder in FOLDERS:
        directory = source / folder
        if not directory.is_dir() or directory.is_symlink() or (hasattr(directory, 'is_junction') and directory.is_junction()):
            raise ValueError(f'Unavailable/linked source folder: {folder}')
        for path in directory.rglob('*'):
            if path.is_symlink() or (hasattr(path, 'is_junction') and path.is_junction()):
                raise ValueError(f'Linked source rejected: {path}')
            if path.is_file():
                resolved = path.resolve(strict=True)
                if source not in resolved.parents:
                    raise ValueError('Source escaped named directory tree')
                paths.append(path)
    paths.sort(key=lambda path: path.relative_to(source).as_posix())
    manifest = {'schema': 'local-research-materials/v1', 'source_root': str(source), 'folders': list(FOLDERS), 'classification': 'user-provided-source-evidence', 'expected_files': args.expected_files, 'files': len(paths), 'source_write_policy': 'read-only', 'extraction_claim': 'Machine text extraction only; images and blocked documents are not claimed read.', 'entries': entries}
    for index, path in enumerate(paths, 1):
        started = time.monotonic()
        entry = {'id': f'material-{index:03}', 'source_relative': path.relative_to(source).as_posix(), 'source_bytes': path.stat().st_size, 'extension': path.suffix.lower(), 'sha256': digest(path)}
        entries.append(entry)
        if entry['sha256'] in by_hash:
            original = by_hash[entry['sha256']]
            entry.update({'status': 'duplicate', 'duplicate_of': original['id'], 'content_status': original['status'], 'text_path': original.get('text_path'), 'format': original.get('format')})
        else:
            by_hash[entry['sha256']] = entry
            try:
                if entry['source_bytes'] > MAX_SOURCE_BYTES:
                    raise ValueError('source_size_limit')
                detected, (records, detail, warnings) = read_one(path)
                text_path = Path('text') / (entry['sha256'] + '.jsonl')
                with (out / text_path).open('w', encoding='utf-8', newline='\n') as stream:
                    for record in records:
                        stream.write(json.dumps(record, ensure_ascii=False, default=str) + '\n')
                entry.update({'format': detected, 'status': 'inventory_only' if detected == 'zip' else 'extracted' if records else 'no_text', 'records': len(records), 'text_chars': sum(len(record['text']) for record in records), 'text_path': text_path.as_posix(), 'text_sha256': digest(out / text_path), 'detail': detail, 'warnings': warnings})
            except (OSError, ValueError, RuntimeError, subprocess.TimeoutExpired, zipfile.BadZipFile, ET.ParseError) as error:
                entry.update({'status': 'blocked' if isinstance(error, RuntimeError) else 'error', 'error_type': type(error).__name__, 'error': str(error)})
            except Exception as error:
                entry.update({'status': 'error', 'error_type': type(error).__name__, 'error': str(error)})
        entry['elapsed_seconds'] = round(time.monotonic() - started, 3)
        manifest['unique_hashes'] = len(by_hash)
        manifest['counts'] = {status: sum(e['status'] == status for e in entries) for status in sorted({e['status'] for e in entries})}
        manifest['indexed_files'] = index
        manifest['complete'] = index == len(paths)
        atomic_json(out / 'manifest.json', manifest)
        if index % 10 == 0 or entry['status'] in ('error', 'blocked'):
            print(json.dumps({'index': index, 'total': len(paths), 'id': entry['id'], 'status': entry['status'], 'source': entry['source_relative']}, ensure_ascii=False), flush=True)
    manifest['count_matches_expected'] = len(paths) == args.expected_files
    manifest['finished_at_utc'] = time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
    atomic_json(out / 'manifest.json', manifest)
    (out / 'manifest.sha256').write_text(digest(out / 'manifest.json') + '  manifest.json\n', encoding='ascii')
    print(json.dumps({'complete': True, 'files': len(paths), 'unique_hashes': len(by_hash), 'counts': manifest['counts'], 'count_matches_expected': manifest['count_matches_expected'], 'manifest': str(out / 'manifest.json')}, ensure_ascii=False))
    return 0 if manifest['count_matches_expected'] else 2

if __name__ == '__main__':
    sys.exit(main())
